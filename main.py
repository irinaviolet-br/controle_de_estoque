"""Labpat inventory control application.

This module implements a Tkinter desktop application that acts as a front-end for
Excel inventory spreadsheets. It supports two user roles: a public operator and a
manager unlocked via a PIN. Both manual and barcode-assisted movements are
available. The manager can also register entries and create new products.
"""
from __future__ import annotations

import base64
import datetime as _dt
import hashlib
import io
import json
import os
import secrets
import smtplib
import ssl
import subprocess
import zipfile
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import tkinter as tk
from tkinter import messagebox, simpledialog, ttk

from email.message import EmailMessage
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

try:
    import msoffcrypto
except ImportError:  # pragma: no cover - optional dependency
    msoffcrypto = None

try:
    from cryptography.fernet import Fernet, InvalidToken
except ImportError:  # pragma: no cover - optional dependency
    Fernet = None  # type: ignore
    InvalidToken = Exception  # type: ignore

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent

APP_TITLE = "Controle de Estoque Labpat"
DATE_FORMAT = "%d/%m/%Y"

STOCK_HEADERS = [
    "CÓDIGO DE BARRAS",
    "FORNECEDOR",
    "PRODUTO",
    "MEDIDA",
    "MODO",
    "DATA DE REPOSIÇÃO",
    "ESTOQUE ATUAL",
    "ESTOQUE CRÍTICO",
]
STOCK_HEADER_ROW = 6
STOCK_DATA_START_ROW = STOCK_HEADER_ROW + 1

LOG_HEADERS = [
    "DATA E HORA",
    "AÇÃO",
    "PRODUTO",
    "ENTRADA",
    "SALDO ANTERIOR",
    "SALDO ATUAL",
    "ESTOQUE",
    "LINHA",
]

# Folder where Excel files are stored. Users can change this path if needed.
DATA_FOLDER = BASE_DIR / "planilhas"

LOGO_FILE = BASE_DIR / "LOGO.png"

SETTINGS_FILE = BASE_DIR / "settings.json"

OWNER_RECOVERY_EMAILS = [
    "fabi.sant.souza@gmail.com",
    "fabiana.santana.souza@outlook.com",
]

DEFAULT_STOCK_NAMES = [
    "ESTOQUE - ESCRITÓRIO",
    "ESTOQUE - TÉCNICA E MACRO",
    "ESTOQUE - COPA",
    "ESTOQUE - HIGIENIZAÇÃO",
    "ESTOQUE - IMUNO",
]

LOG_FILE = DATA_FOLDER / "LOG.xlsx"

# ---------------------------------------------------------------------------
# Settings management
# ---------------------------------------------------------------------------


class AppSettings:
    """Persist application-level settings such as PINs and email configuration."""

    DEFAULTS = {
        "manager_pin": "2468",
        "manager_email": "qualidade@labpat.med.br",
        "stock_password": "789",
        "email": {
            "host": "",
            "port": 587,
            "username": "",
            "password": "",
            "use_tls": True,
            "use_ssl": False,
            "from_address": "",
        },
        "owner_email": {
            "host": "",
            "port": 587,
            "username": "",
            "password": "",
            "use_tls": True,
            "use_ssl": False,
            "from_address": "",
        },
    }

    def __init__(self, path: Path = SETTINGS_FILE):
        if Fernet is None:
            raise RuntimeError(
                "Instale o pacote 'cryptography' para executar o aplicativo de controle de estoque."
            )
        self.path = path
        self._data = json.loads(json.dumps(self.DEFAULTS))
        self._fernet = self._build_fernet()
        self._ensure_file_exists()
        self.load()

    # -- persistence ----------------------------------------------------

    def load(self) -> None:
        if not self.path.exists():
            return
        try:
            raw = self.path.read_bytes()
        except OSError:
            return
        if not raw:
            return

        try:
            decrypted = self._fernet.decrypt(raw)
            payload = json.loads(decrypted.decode("utf-8"))
        except InvalidToken:
            try:
                payload = json.loads(raw.decode("utf-8"))
            except (json.JSONDecodeError, UnicodeDecodeError):
                return
            else:
                self._deep_update(self._data, payload)
                self.save(silent=True)
                return
        except (json.JSONDecodeError, UnicodeDecodeError):
            return

        self._deep_update(self._data, payload)

    def save(self, *, silent: bool = False) -> None:
        try:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            payload = json.dumps(self._data, indent=2, ensure_ascii=False).encode("utf-8")
            encrypted = self._fernet.encrypt(payload)
            with self.path.open("wb") as fh:
                fh.write(encrypted)
        except OSError:
            if not silent:
                messagebox.showwarning(
                    APP_TITLE,
                    "Não foi possível salvar as configurações em settings.json.",
                )

    # -- helpers --------------------------------------------------------

    def _deep_update(self, target: dict, updates: dict) -> None:
        for key, value in updates.items():
            if isinstance(value, dict) and isinstance(target.get(key), dict):
                self._deep_update(target[key], value)
            else:
                target[key] = value

    def _build_fernet(self) -> "Fernet":
        secret = hashlib.sha256(b"labpat-inventory-settings-key-v1").digest()
        key = base64.urlsafe_b64encode(secret)
        return Fernet(key)

    # -- properties -----------------------------------------------------

    @property
    def manager_pin(self) -> str:
        raw = str(self._data.get("manager_pin", self.DEFAULTS["manager_pin"]))
        return raw.strip()

    @property
    def manager_email(self) -> str:
        return str(self._data.get("manager_email", self.DEFAULTS["manager_email"])).strip()

    @property
    def stock_password(self) -> str:
        return str(self._data.get("stock_password", self.DEFAULTS["stock_password"]))

    @property
    def email_settings(self) -> dict:
        return dict(self._data.get("email", {}))

    @property
    def owner_email_settings(self) -> dict:
        return dict(self._data.get("owner_email", {}))

    # -- mutators -------------------------------------------------------

    def update_manager_credentials(self, pin: str, email: str) -> None:
        self._data["manager_pin"] = pin
        self._data["manager_email"] = email
        self.save()

    def update_stock_password(self, password: str) -> None:
        self._data["stock_password"] = password
        self.save()

    # -- internal -------------------------------------------------------

    def _ensure_file_exists(self) -> None:
        if self.path.exists():
            return
        try:
            self.save(silent=True)
        except Exception:
            # A gravação inicial é um best-effort; qualquer falha é ignorada
            # aqui e será tratada quando a aplicação tentar salvar novamente.
            pass


# ---------------------------------------------------------------------------
# Email delivery helpers
# ---------------------------------------------------------------------------


class EmailService:
    def __init__(self, settings: AppSettings):
        self.settings = settings

    def send_manager_pin(
        self,
        pin: str,
        subject: str,
        reason: str,
        target_email: Optional[str] = None,
    ) -> None:
        manager_email = (target_email or self.settings.manager_email).strip()
        if not manager_email:
            raise ValueError("Nenhum e-mail do gestor cadastrado.")

        subject = subject.strip() or "Controle de Estoque - PIN"

        def build_message(target: str, from_address: str) -> EmailMessage:
            message = EmailMessage()
            message["Subject"] = subject
            message["From"] = from_address
            message["To"] = target
            message.set_content(
                "Olá,\n\n"
                f"{reason}\n\n"
                f"PIN de acesso: {pin}\n\n"
                "Esta mensagem foi enviada automaticamente pelo sistema de estoque."
            )
            return message

        manager_config = self.settings.email_settings
        manager_from = self._resolve_from_address(manager_config)
        manager_message = build_message(manager_email, manager_from)
        self._send_messages(
            manager_config,
            manager_from,
            [(manager_email, manager_message)],
            missing_host_error=(
                "Configure o servidor SMTP no arquivo settings.json (seção \"email\") para enviar e-mails de recuperação."
            ),
        )

        owner_recipients = [email for email in OWNER_RECOVERY_EMAILS if email != manager_email]
        if not owner_recipients:
            return

        owner_config = self.settings.owner_email_settings
        owner_from = self._resolve_from_address(owner_config)
        owner_messages = [
            (recipient, build_message(recipient, owner_from)) for recipient in owner_recipients
        ]
        self._send_messages(
            owner_config,
            owner_from,
            owner_messages,
            missing_host_error=(
                "Configure o servidor SMTP no arquivo settings.json (seção \"owner_email\") para enviar as cópias de recuperação."
            ),
        )

    def _resolve_from_address(self, config: dict) -> str:
        return (
            config.get("from_address")
            or config.get("username")
            or "no-reply@localhost"
        )

    def _send_messages(
        self,
        config: dict,
        from_address: str,
        messages: Iterable[Tuple[str, EmailMessage]],
        *,
        missing_host_error: str,
    ) -> None:
        host = config.get("host")
        if not host:
            raise ValueError(missing_host_error)

        port = int(config.get("port") or (465 if config.get("use_ssl") else 587))
        username = config.get("username") or None
        password = config.get("password") or None
        use_ssl = bool(config.get("use_ssl"))
        use_tls = bool(config.get("use_tls", True)) and not use_ssl

        try:
            if use_ssl:
                context = ssl.create_default_context()
                with smtplib.SMTP_SSL(host, port, context=context) as server:
                    if username and password:
                        server.login(username, password)
                    for recipient, message in messages:
                        server.sendmail(from_address, [recipient], message.as_string())
            else:
                with smtplib.SMTP(host, port) as server:
                    if use_tls:
                        server.starttls(context=ssl.create_default_context())
                    if username and password:
                        server.login(username, password)
                    for recipient, message in messages:
                        server.sendmail(from_address, [recipient], message.as_string())
        except OSError as exc:
            raise ValueError(f"Falha ao enviar o e-mail: {exc}") from exc

# ---------------------------------------------------------------------------
# Utilities
# ---------------------------------------------------------------------------


def ensure_directory(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def discover_stock_files(folder: Path = DATA_FOLDER) -> Dict[str, Path]:
    """Return a mapping of stock names to workbook paths found on disk."""

    ensure_directory(folder)
    files: Dict[str, Path] = {}

    for path in sorted(folder.glob("ESTOQUE*.xlsx")):
        if path.is_file():
            files[path.stem] = path

    for name in DEFAULT_STOCK_NAMES:
        files.setdefault(name, folder / f"{name}.xlsx")

    return files


def normalise_product_name(name: str) -> str:
    return name.strip().lower()


def parse_positive_number(raw: str) -> Optional[float]:
    try:
        value = float(raw.replace(",", "."))
    except (ValueError, AttributeError):
        return None
    if value <= 0:
        return None
    return value


def open_file_with_default_app(path: Path) -> None:
    path = path.resolve()
    if not path.exists():
        raise ValueError(f"Arquivo não encontrado: {path}")
    try:
        if sys.platform.startswith("win"):
            os.startfile(str(path))  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)])
    except Exception as exc:
        raise ValueError(f"Não foi possível abrir a planilha selecionada: {exc}") from exc


# ---------------------------------------------------------------------------
# Data access layer
# ---------------------------------------------------------------------------


@dataclass
class ProductRecord:
    row: int
    barcode: str
    supplier: str
    name: str
    measure: str
    mode: str
    restock_date: Optional[_dt.date]
    stock: float
    critical_stock: float


class InventoryWorkbook:
    """Wraps an Excel workbook and exposes helpers for accessing stock data."""

    def __init__(self, path: Path, settings: AppSettings):
        self.path = path
        self.settings = settings
        ensure_directory(path.parent)
        self._ensure_template()
        # Garantir que arquivos existentes estejam protegidos com senha.
        with self._load(save_on_exit=True):
            pass

    # -- public helpers --------------------------------------------------

    def ensure_exists(self) -> Path:
        if not self.path.exists():
            with self._load(save_on_exit=True):
                pass
        return self.path

    def list_products(self) -> List[str]:
        with self._load() as ws:
            return [record.name for record in self._iter_products(ws)]

    def get_product_by_name(self, name: str) -> Optional[ProductRecord]:
        normalised = normalise_product_name(name)
        with self._load() as ws:
            for record in self._iter_products(ws):
                if normalise_product_name(record.name) == normalised:
                    return record
        return None

    def get_product_by_barcode(self, barcode: str) -> Optional[ProductRecord]:
        barcode = (barcode or "").strip()
        if not barcode:
            return None
        with self._load() as ws:
            for record in self._iter_products(ws):
                if str(record.barcode).strip() == barcode:
                    return record
        return None

    def update_stock(
        self,
        row: int,
        delta: float,
        restock_date: Optional[_dt.date],
    ) -> Tuple[float, float]:
        """Applies a delta to the stock of the given row and returns before/after."""
        with self._load(save_on_exit=True) as ws:
            current = float(ws.cell(row=row, column=7).value or 0)
            updated = current + delta
            if updated < 0:
                raise ValueError("Saldo insuficiente para a operação solicitada.")
            ws.cell(row=row, column=7, value=updated)
            if restock_date:
                ws.cell(row=row, column=6, value=restock_date.strftime(DATE_FORMAT))
            return current, updated

    def add_product(
        self,
        barcode: str,
        name: str,
        supplier: str,
        measure: str,
        mode: str,
        critical_stock: float,
    ) -> int:
        with self._load(save_on_exit=True) as ws:
            row = max(ws.max_row + 1, STOCK_DATA_START_ROW)
            ws.cell(row=row, column=1, value=barcode.strip())
            ws.cell(row=row, column=2, value=supplier.strip())
            ws.cell(row=row, column=3, value=name.strip())
            ws.cell(row=row, column=4, value=measure.strip())
            ws.cell(row=row, column=5, value=mode.strip().upper())
            ws.cell(row=row, column=6, value=_dt.date.today().strftime(DATE_FORMAT))
            ws.cell(row=row, column=7, value=0)
            ws.cell(row=row, column=8, value=critical_stock)
            return row

    # -- context manager helpers ----------------------------------------

    class _WorksheetManager:
        def __init__(self, workbook: "InventoryWorkbook", save_on_exit: bool):
            self.workbook = workbook
            self.save_on_exit = save_on_exit
            self.wb = None
            self.ws = None

        def __enter__(self) -> Worksheet:
            if self.workbook.path.exists():
                try:
                    self.wb = load_workbook(self.workbook.path)
                    self._encrypted = False
                except (InvalidFileException, zipfile.BadZipFile):
                    if msoffcrypto is None:
                        raise ValueError(
                            "Instale o pacote msoffcrypto-tool para abrir planilhas protegidas com senha."
                        )
                    self._encrypted = True
                    decrypted = io.BytesIO()
                    try:
                        with self.workbook.path.open("rb") as fh:
                            office = msoffcrypto.OfficeFile(fh)
                            office.load_key(password=self.workbook.settings.stock_password)
                            office.decrypt(decrypted)
                    except Exception as exc:  # pragma: no cover - depends on external lib
                        raise ValueError(
                            "Não foi possível abrir a planilha protegida. Verifique a senha informada."
                        ) from exc
                    decrypted.seek(0)
                    self.wb = load_workbook(decrypted)
                    self._decrypted_stream = decrypted
            else:
                self.wb = Workbook()
                self._encrypted = False
            sheet_name = self.workbook._sheet_name
            if sheet_name in self.wb.sheetnames:
                self.ws = self.wb[sheet_name]
            else:
                self.ws = self.wb.active
                self.ws.title = sheet_name
            self.workbook._ensure_headers(self.ws)
            self.workbook._apply_protection(self.ws)
            return self.ws

        def __exit__(self, exc_type, exc_val, exc_tb) -> None:
            if self.save_on_exit and exc_type is None:
                self.workbook._apply_protection(self.ws)
                if getattr(self, "_encrypted", False) and msoffcrypto is not None:
                    buffer = io.BytesIO()
                    self.wb.save(buffer)
                    buffer.seek(0)
                    if hasattr(msoffcrypto.OfficeFile, "encrypt"):
                        with self.workbook.path.open("wb") as fh:  # pragma: no cover - I/O
                            office = msoffcrypto.OfficeFile(buffer)
                            office.encrypt(
                                password=self.workbook.settings.stock_password,
                                output=fh,
                            )
                    else:  # pragma: no cover - depends on optional lib
                        self.wb.save(self.workbook.path)
                else:
                    self.wb.save(self.workbook.path)
            elif not self.save_on_exit:
                self.wb.close()
            else:
                # Discard changes on error.
                self.wb.close()

    def _load(self, save_on_exit: bool = False) -> "InventoryWorkbook._WorksheetManager":
        return InventoryWorkbook._WorksheetManager(self, save_on_exit)

    # -- private helpers -------------------------------------------------

    @property
    def _sheet_name(self) -> str:
        return self.path.stem

    def _ensure_template(self) -> None:
        if not self.path.exists():
            wb = Workbook()
            ws = wb.active
            ws.title = self._sheet_name
            self._ensure_headers(ws)
            self._apply_protection(ws)
            wb.save(self.path)

    def refresh_password(self) -> None:
        with self._load(save_on_exit=True):
            pass

    def _ensure_headers(self, ws: Worksheet) -> None:
        for col in range(1, len(STOCK_HEADERS) + 1):
            ws.cell(row=STOCK_HEADER_ROW, column=col, value=STOCK_HEADERS[col - 1])

    def _apply_protection(self, ws: Worksheet) -> None:
        # Lock the worksheet so it cannot be edited directly in Excel without a password.
        ws.protection.sheet = True
        ws.protection.set_password(self.settings.stock_password)

    def _iter_products(self, ws: Worksheet) -> Iterable[ProductRecord]:
        for row in range(STOCK_DATA_START_ROW, ws.max_row + 1):
            name = ws.cell(row=row, column=3).value
            if not name:
                continue
            barcode = ws.cell(row=row, column=1).value or ""
            supplier = ws.cell(row=row, column=2).value or ""
            measure = ws.cell(row=row, column=4).value or ""
            mode = ws.cell(row=row, column=5).value or ""
            restock_date_raw = ws.cell(row=row, column=6).value
            if isinstance(restock_date_raw, _dt.datetime):
                restock_date = restock_date_raw.date()
            elif isinstance(restock_date_raw, _dt.date):
                restock_date = restock_date_raw
            elif isinstance(restock_date_raw, str):
                try:
                    restock_date = _dt.datetime.strptime(restock_date_raw, DATE_FORMAT).date()
                except ValueError:
                    restock_date = None
            else:
                restock_date = None
            stock = float(ws.cell(row=row, column=7).value or 0)
            critical_stock = float(ws.cell(row=row, column=8).value or 0)
            yield ProductRecord(
                row=row,
                barcode=str(barcode),
                supplier=str(supplier),
                name=str(name),
                measure=str(measure),
                mode=str(mode),
                restock_date=restock_date,
                stock=stock,
                critical_stock=critical_stock,
            )


class LogWorkbook:
    def __init__(self, path: Path):
        self.path = path
        ensure_directory(path.parent)
        self._ensure_template()

    def append(
        self,
        action: str,
        product: str,
        delta: float,
        previous: float,
        updated: float,
        stock_name: str,
        row: int,
    ) -> None:
        wb = load_workbook(self.path)
        ws = wb.active
        timestamp = _dt.datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        entry = [
            timestamp,
            action.upper(),
            product,
            delta,
            previous,
            updated,
            stock_name,
            row,
        ]
        ws.append(entry)
        wb.save(self.path)

    def _ensure_template(self) -> None:
        if not self.path.exists():
            wb = Workbook()
            ws = wb.active
            ws.title = "LOG"
            for idx, header in enumerate(LOG_HEADERS, start=1):
                ws.cell(row=1, column=idx, value=header)
            wb.save(self.path)

    def ensure_exists(self) -> Path:
        self._ensure_template()
        return self.path


# ---------------------------------------------------------------------------
# Application services
# ---------------------------------------------------------------------------


class InventoryService:
    def __init__(self, settings: AppSettings, email_service: EmailService):
        self.settings = settings
        self.email_service = email_service
        stock_files = discover_stock_files()
        self.workbooks = {
            name: InventoryWorkbook(path, settings) for name, path in stock_files.items()
        }
        self.log = LogWorkbook(LOG_FILE)

    def get_workbook(self, name: str) -> InventoryWorkbook:
        return self.workbooks[name]

    def list_stock_names(self) -> List[str]:
        return list(self.workbooks.keys())

    def list_all_planilhas(self) -> List[str]:
        names = list(self.workbooks.keys())
        if "LOG" not in names:
            names.append("LOG")
        return names

    def get_workbook_path(self, name: str) -> Path:
        if name == "LOG":
            return self.log.ensure_exists()
        workbook = self.workbooks.get(name)
        if workbook is None:
            raise ValueError("Planilha não encontrada.")
        return workbook.ensure_exists()

    def get_manager_email(self) -> str:
        return self.settings.manager_email

    def get_sheet_preview(
        self, name: str, start_row: int = 5
    ) -> Tuple[List[str], List[Tuple[List[str], str]]]:
        path = self.get_workbook_path(name)
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except (InvalidFileException, OSError) as exc:
            raise ValueError(f"Não foi possível abrir a planilha selecionada: {exc}") from exc
        ws = wb.active
        header_row = 1 if name == "LOG" else STOCK_HEADER_ROW
        data_start = max(start_row, header_row + 1)

        max_col = 0
        for row in ws.iter_rows(min_row=header_row, max_row=ws.max_row, values_only=True):
            for idx, value in enumerate(row, start=1):
                if value is None:
                    continue
                if isinstance(value, str) and not value.strip():
                    continue
                max_col = max(max_col, idx)

        if max_col == 0:
            return [], []

        columns = [get_column_letter(idx) for idx in range(1, max_col + 1)]
        formatted_rows: List[Tuple[List[str], str]] = []

        header_values = [
            self._format_cell(ws.cell(row=header_row, column=idx).value)
            for idx in range(1, max_col + 1)
        ]
        formatted_rows.append((header_values, "headers"))

        data_rows = list(
            ws.iter_rows(min_row=data_start, max_row=ws.max_row, values_only=True)
        )
        data_index = 0
        for row in data_rows:
            values = [
                self._format_cell(row[idx]) if idx < len(row) else ""
                for idx in range(max_col)
            ]
            if not any(value.strip() for value in values):
                continue
            tag = "even" if data_index % 2 == 0 else "odd"
            data_index += 1
            formatted_rows.append((values, tag))
        wb.close()
        return columns, formatted_rows

    def _format_cell(self, value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, _dt.datetime):
            return value.strftime("%d/%m/%Y %H:%M:%S")
        if isinstance(value, _dt.date):
            return value.strftime(DATE_FORMAT)
        return str(value)

    # -- authentication -------------------------------------------------

    def verify_manager_pin(self, pin: str) -> bool:
        return pin == self.settings.manager_pin

    def update_manager_pin(self, current_pin: str, new_pin: str, email: str) -> None:
        if current_pin != self.settings.manager_pin:
            raise ValueError("Senha atual incorreta.")
        self._validate_pin(new_pin)
        email = email.strip()
        if not email:
            raise ValueError("Informe o e-mail do gestor.")
        self.email_service.send_manager_pin(
            new_pin,
            subject="Nova senha de acesso ao estoque",
            reason="A senha do gestor foi atualizada manualmente.",
            target_email=email,
        )
        self.settings.update_manager_credentials(new_pin, email)

    def recover_manager_pin(self) -> str:
        current_pin = self.settings.manager_pin
        self._validate_pin(current_pin)
        self.email_service.send_manager_pin(
            current_pin,
            subject="Recuperação de senha do estoque",
            reason="A senha atual foi enviada conforme solicitado no recurso de recuperação.",
        )
        return current_pin

    def _generate_random_pin(self) -> str:
        digits = "0123456789"
        length = 6
        return "".join(secrets.choice(digits) for _ in range(length))

    def _validate_pin(self, pin: str) -> None:
        if not pin.isdigit():
            raise ValueError("A senha deve conter apenas números.")
        if not (1 <= len(pin) <= 10):
            raise ValueError("A senha deve ter entre 1 e 10 dígitos.")

    # -- stock password -------------------------------------------------

    def update_stock_password(self, password: str) -> None:
        password = password.strip()
        if not password:
            raise ValueError("Informe uma senha válida para as planilhas.")
        self.settings.update_stock_password(password)
        for workbook in self.workbooks.values():
            workbook.refresh_password()

    def move_stock(
        self,
        stock_name: str,
        product: ProductRecord,
        delta: float,
        restock_date: Optional[_dt.date],
    ) -> Tuple[float, float]:
        workbook = self.get_workbook(stock_name)
        before, after = workbook.update_stock(product.row, delta, restock_date)
        action = "ENTRADA" if delta > 0 else "SAÍDA"
        self.log.append(action, product.name, delta, before, after, stock_name, product.row)
        return before, after

    def add_product(
        self,
        stock_name: str,
        barcode: str,
        name: str,
        supplier: str,
        measure: str,
        mode: str,
        critical_stock: float,
    ) -> int:
        workbook = self.get_workbook(stock_name)
        row = workbook.add_product(barcode, name, supplier, measure, mode, critical_stock)
        self.log.append(
            "NOVO PRODUTO",
            name,
            0,
            0,
            0,
            stock_name,
            row,
        )
        return row


# ---------------------------------------------------------------------------
# GUI components
# ---------------------------------------------------------------------------


class WorkbookSelector(simpledialog.Dialog):
    def __init__(
        self,
        parent: tk.Misc,
        choices: Iterable[str],
        *,
        title: str = "Saída – escolha o setor",
        prompt: str = "Escolha o estoque:",
    ):
        self.choices = list(choices)
        self.selection: Optional[str] = None
        self.prompt = prompt
        super().__init__(parent, title=title)

    def body(self, master):
        ttk.Label(master, text=self.prompt).grid(row=0, column=0, padx=10, pady=(10, 5))
        self.combo = ttk.Combobox(master, values=self.choices, state="readonly")
        self.combo.grid(row=1, column=0, padx=10, pady=(0, 10))
        if self.choices:
            self.combo.current(0)
        return self.combo

    def apply(self):
        self.selection = self.combo.get()


class AutoCompleteCombobox(ttk.Combobox):
    """Combobox that filters values as the user types."""

    def __init__(self, master=None, **kwargs):
        values = list(kwargs.pop("values", []))
        super().__init__(master, values=values, **kwargs)
        self._all_values = values
        self._filtered_values = values
        self.configure(state="normal")
        self.bind("<KeyRelease>", self._on_keyrelease, add="+")
        self.bind("<FocusIn>", self._restore_all, add="+")

    def set_completion_list(self, values: Iterable[str]) -> None:
        self._all_values = list(values)
        self._filtered_values = self._all_values
        self.configure(values=self._all_values)

    def first_match(self) -> Optional[str]:
        return self._filtered_values[0] if self._filtered_values else None

    def _restore_all(self, event=None):
        if not self.get():
            self.configure(values=self._all_values)
            self._filtered_values = self._all_values

    def _on_keyrelease(self, event):
        # Ignore navigation keys so the default Combobox behaviour stays intact.
        if event.keysym in {"Left", "Right", "Up", "Down", "Home", "End", "Tab"}:
            return
        if event.keysym == "Escape":
            self.delete(0, tk.END)
            self.configure(values=self._all_values)
            self._filtered_values = self._all_values
            return

        pattern = normalise_product_name(self.get())
        if not pattern:
            filtered = self._all_values
        else:
            filtered = [
                item
                for item in self._all_values
                if normalise_product_name(item).startswith(pattern)
            ]
        if filtered:
            self.configure(values=filtered)
            self._filtered_values = filtered
        else:
            self.configure(values=self._all_values)
            self._filtered_values = self._all_values

        if event.keysym == "Return":
            if self._filtered_values:
                self.set(self._filtered_values[0])
                self.icursor(tk.END)
            # When pressing Enter with text typed, trigger selection handlers.
            self.event_generate("<<ComboboxSelected>>")


class ManualMovementWindow(tk.Toplevel):
    def __init__(
        self,
        master: tk.Tk,
        service: InventoryService,
        stock_name: str,
        is_entry: bool,
        manager_mode: bool,
    ):
        super().__init__(master)
        self.service = service
        self.stock_name = stock_name
        self.is_entry = is_entry
        self.manager_mode = manager_mode
        self.product: Optional[ProductRecord] = None

        self.title("Movimento de Estoque")
        self.resizable(False, False)

        ttk.Label(self, text=f"Movimento: {'Entrada' if is_entry else 'Saída'}").grid(
            row=0, column=0, columnspan=2, padx=10, pady=(10, 0), sticky="w"
        )

        ttk.Label(self, text="Produto").grid(row=1, column=0, padx=10, pady=(10, 0), sticky="w")
        workbook = self.service.get_workbook(stock_name)
        try:
            products = workbook.list_products()
        except ValueError as exc:
            messagebox.showerror(APP_TITLE, str(exc))
            self.destroy()
            return
        self.product_combo = AutoCompleteCombobox(self, values=products, width=60)
        self.product_combo.grid(row=2, column=0, columnspan=2, padx=10, pady=(0, 5), sticky="we")
        self.product_combo.bind("<<ComboboxSelected>>", self._on_product_selected)
        self.product_combo.bind("<FocusOut>", self._on_product_selected, add="+")

        self.stock_label = ttk.Label(self, text="Estoque atual: -")
        self.stock_label.grid(row=3, column=0, columnspan=2, padx=10, pady=(0, 5), sticky="w")

        ttk.Label(self, text="Quantidade").grid(row=4, column=0, padx=10, pady=(0, 0), sticky="w")
        self.quantity_entry = ttk.Entry(self)
        self.quantity_entry.grid(row=5, column=0, columnspan=2, padx=10, pady=(0, 10), sticky="we")

        if self.is_entry and self.manager_mode:
            self.restock_var = tk.BooleanVar(value=False)
            self.date_var = tk.StringVar()
            self.date_checkbox = ttk.Checkbutton(
                self,
                text="Informar/alterar data de reposição caso não tenha sido hoje",
                variable=self.restock_var,
                command=self._toggle_date_entry,
            )
            self.date_checkbox.grid(row=6, column=0, columnspan=2, padx=10, pady=(0, 5), sticky="w")

            self.date_entry = ttk.Entry(self, textvariable=self.date_var, state="disabled")
            self.date_entry.grid(row=7, column=0, columnspan=2, padx=10, pady=(0, 10), sticky="we")
        else:
            self.restock_var = None
            self.date_var = None
            self.date_entry = None

        btn_frame = ttk.Frame(self)
        btn_frame.grid(row=8, column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame, text="Confirmar", command=self._confirm).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="Fechar", command=self.destroy).grid(row=0, column=1, padx=5)

    def _toggle_date_entry(self):
        if self.date_entry is None:
            return
        if self.restock_var.get():
            self.date_entry.configure(state="normal")
            if not self.date_var.get():
                self.date_var.set(_dt.date.today().strftime(DATE_FORMAT))
        else:
            self.date_entry.configure(state="disabled")
            self.date_var.set("")

    def _on_product_selected(self, event=None):
        name = self.product_combo.get().strip()
        if not name and isinstance(self.product_combo, AutoCompleteCombobox):
            match = self.product_combo.first_match()
            if match:
                self.product_combo.set(match)
                name = match
        if not name:
            return
        workbook = self.service.get_workbook(self.stock_name)
        try:
            self.product = workbook.get_product_by_name(name)
        except ValueError as exc:
            messagebox.showerror(APP_TITLE, str(exc))
            self.destroy()
            return
        if not self.product and isinstance(self.product_combo, AutoCompleteCombobox):
            match = self.product_combo.first_match()
            if match and match != name:
                self.product_combo.set(match)
                try:
                    self.product = workbook.get_product_by_name(match)
                except ValueError as exc:
                    messagebox.showerror(APP_TITLE, str(exc))
                    self.destroy()
                    return
        if not self.product:
            self.stock_label.configure(text="Produto não encontrado")
            return
        self.stock_label.configure(
            text=f"Estoque atual: {self.product.stock:g} {self.product.mode.upper()}"
        )

    def _confirm(self):
        if not self.product:
            messagebox.showerror(APP_TITLE, "Selecione um produto válido.")
            return
        quantity = parse_positive_number(self.quantity_entry.get())
        if quantity is None:
            messagebox.showerror(APP_TITLE, "Informe uma quantidade numérica maior que zero.")
            return
        delta = quantity if self.is_entry else -quantity

        restock_date = None
        if self.is_entry and self.manager_mode and self.restock_var.get():
            try:
                restock_date = _dt.datetime.strptime(self.date_var.get(), DATE_FORMAT).date()
            except (TypeError, ValueError):
                messagebox.showerror(APP_TITLE, "Data de reposição inválida. Use dd/mm/aaaa.")
                return
        try:
            previous, updated = self.service.move_stock(
                self.stock_name, self.product, delta, restock_date
            )
        except ValueError as exc:
            messagebox.showerror(APP_TITLE, str(exc))
            return
        messagebox.showinfo(
            APP_TITLE,
            f"Operação bem-sucedida.\n\nSaldo anterior: {previous:g}\nSaldo atual: {updated:g}",
            parent=self,
        )
        if self.product:
            self.product.stock = updated
            self.stock_label.configure(
                text=f"Estoque atual: {updated:g} {self.product.mode.upper()}"
            )
        self.quantity_entry.delete(0, tk.END)
        self.quantity_entry.focus_set()


class ScannerWindow(tk.Toplevel):
    def __init__(
        self,
        master: tk.Tk,
        service: InventoryService,
        stock_name: str,
        manager_mode: bool,
    ):
        super().__init__(master)
        self.service = service
        self.stock_name = stock_name
        self.manager_mode = manager_mode
        self.product: Optional[ProductRecord] = None

        self.title("Leitor de almoxarifado")
        self.resizable(False, False)

        ttk.Label(self, text="Código de barras:").grid(row=0, column=0, padx=10, pady=(10, 0), sticky="w")
        self.barcode_var = tk.StringVar()
        self.barcode_entry = ttk.Entry(self, textvariable=self.barcode_var, width=40)
        self.barcode_entry.grid(row=1, column=0, columnspan=2, padx=10, pady=(0, 5), sticky="we")
        self.barcode_entry.bind("<Return>", self._lookup_barcode)

        ttk.Label(self, text="Estoque:").grid(row=2, column=0, padx=10, pady=(5, 0), sticky="w")
        self.product_label = ttk.Label(self, text="-", width=60)
        self.product_label.grid(row=3, column=0, columnspan=2, padx=10, sticky="w")

        ttk.Label(self, text="Quantidade:").grid(row=4, column=0, padx=10, pady=(10, 0), sticky="w")
        self.quantity_entry = ttk.Entry(self)
        self.quantity_entry.grid(row=5, column=0, columnspan=2, padx=10, pady=(0, 10), sticky="we")

        self.mode_var = tk.StringVar(value="saida")
        if self.manager_mode:
            ttk.Radiobutton(
                self, text="Saída", variable=self.mode_var, value="saida"
            ).grid(row=6, column=0, padx=10, sticky="w")
            ttk.Radiobutton(
                self, text="Entrada", variable=self.mode_var, value="entrada"
            ).grid(row=6, column=1, padx=10, sticky="w")

            ttk.Label(self, text="Data de reposição (dd/mm/aaaa):").grid(
                row=7, column=0, columnspan=2, padx=10, pady=(5, 0), sticky="w"
            )
            self.date_var = tk.StringVar()
            self.date_entry = ttk.Entry(self, textvariable=self.date_var)
            self.date_entry.grid(row=8, column=0, columnspan=2, padx=10, pady=(0, 10), sticky="we")
        else:
            ttk.Checkbutton(self, text="Saída", variable=self.mode_var, onvalue="saida").grid(
                row=6, column=0, columnspan=2, padx=10, sticky="w"
            )
            self.mode_var.set("saida")
            self.date_var = None
            self.date_entry = None

        btn_frame = ttk.Frame(self)
        btn_frame.grid(row=9, column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame, text="Confirmar", command=self._confirm).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="Fechar", command=self.destroy).grid(row=0, column=1, padx=5)

    def _lookup_barcode(self, event=None):
        barcode = self.barcode_var.get().strip()
        workbook = self.service.get_workbook(self.stock_name)
        try:
            self.product = workbook.get_product_by_barcode(barcode)
        except ValueError as exc:
            messagebox.showerror(APP_TITLE, str(exc))
            self.destroy()
            return
        if not self.product:
            self.product_label.configure(text="Produto não encontrado")
            return
        self.product_label.configure(
            text=f"{self.product.name}\nEstoque atual: {self.product.stock:g} {self.product.mode.upper()}"
        )
        self.quantity_entry.focus_set()

    def _confirm(self):
        if not self.product:
            messagebox.showerror(APP_TITLE, "Informe um código de barras válido e pressione ENTER.")
            return
        quantity = parse_positive_number(self.quantity_entry.get())
        if quantity is None:
            messagebox.showerror(APP_TITLE, "Informe uma quantidade numérica maior que zero.")
            return
        action = self.mode_var.get()
        if action not in {"entrada", "saida"}:
            action = "saida"
        delta = quantity if action == "entrada" else -quantity

        restock_date = None
        if action == "entrada" and self.manager_mode:
            if self.date_var.get():
                try:
                    restock_date = _dt.datetime.strptime(self.date_var.get(), DATE_FORMAT).date()
                except ValueError:
                    messagebox.showerror(APP_TITLE, "Data de reposição inválida. Use dd/mm/aaaa.")
                    return
            else:
                restock_date = _dt.date.today()

        try:
            previous, updated = self.service.move_stock(
                self.stock_name, self.product, delta, restock_date
            )
        except ValueError as exc:
            messagebox.showerror(APP_TITLE, str(exc))
            return
        messagebox.showinfo(
            APP_TITLE,
            f"Operação bem-sucedida.\n\nSaldo anterior: {previous:g}\nSaldo atual: {updated:g}",
            parent=self,
        )
        self.product = None
        self.barcode_var.set("")
        self.product_label.configure(text="-")
        self.quantity_entry.delete(0, tk.END)
        if self.date_var is not None:
            self.date_var.set("")
        self.barcode_entry.focus_set()


class NewProductWindow(tk.Toplevel):
    def __init__(self, master: tk.Tk, service: InventoryService, stock_name: str):
        super().__init__(master)
        self.service = service
        self.stock_name = stock_name
        self.title("Cadastrar novo produto")
        self.resizable(False, False)
        self.columnconfigure(0, weight=1)

        ttk.Label(self, text="Código de Barras").grid(row=0, column=0, padx=10, pady=(10, 0), sticky="w")
        self.barcode_entry = ttk.Entry(self, width=60)
        self.barcode_entry.grid(row=1, column=0, padx=10, pady=(0, 5), sticky="we")

        ttk.Label(self, text="Produto").grid(row=2, column=0, padx=10, pady=(5, 0), sticky="w")
        self.product_entry = ttk.Entry(self, width=60)
        self.product_entry.grid(row=3, column=0, padx=10, pady=(0, 5), sticky="we")

        ttk.Label(self, text="Fornecedor").grid(row=4, column=0, padx=10, pady=(5, 0), sticky="w")
        self.supplier_entry = ttk.Entry(self, width=60)
        self.supplier_entry.grid(row=5, column=0, padx=10, pady=(0, 5), sticky="we")

        ttk.Label(self, text="Estoque crítico").grid(row=6, column=0, padx=10, pady=(5, 0), sticky="w")
        self.critical_entry = ttk.Entry(self, width=60)
        self.critical_entry.insert(0, "0")
        self.critical_entry.grid(row=7, column=0, padx=10, pady=(0, 5), sticky="we")

        ttk.Label(self, text="Modo do estoque").grid(row=8, column=0, padx=10, pady=(5, 0), sticky="w")
        self.mode_combo = ttk.Combobox(self, values=["UN", "ML"], state="readonly", width=57)
        self.mode_combo.current(0)
        self.mode_combo.grid(row=9, column=0, padx=10, pady=(0, 10), sticky="we")

        ttk.Label(self, text="Medida").grid(row=10, column=0, padx=10, pady=(0, 0), sticky="w")
        self.measure_entry = ttk.Entry(self, width=60)
        self.measure_entry.grid(row=11, column=0, padx=10, pady=(0, 10), sticky="we")

        btn_frame = ttk.Frame(self)
        btn_frame.grid(row=12, column=0, pady=10)
        ttk.Button(btn_frame, text="Salvar", command=self._save).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="Cancelar", command=self.destroy).grid(row=0, column=1, padx=5)

        self.update_idletasks()
        width = self.winfo_reqwidth() * 3
        height = self.winfo_reqheight()
        self.geometry(f"{width}x{height}")
        self.minsize(width, height)

    def _save(self):
        barcode = self.barcode_entry.get().strip()
        product = self.product_entry.get().strip()
        supplier = self.supplier_entry.get().strip()
        measure = self.measure_entry.get().strip()
        mode = self.mode_combo.get().strip() or "UN"
        critical_raw = self.critical_entry.get().strip() or "0"
        try:
            critical = float(critical_raw.replace(",", "."))
            if critical < 0:
                raise ValueError
        except ValueError:
            messagebox.showerror(APP_TITLE, "Estoque crítico deve ser um número maior ou igual a zero.")
            return

        if not product:
            messagebox.showerror(APP_TITLE, "Informe o nome do produto.")
            return

        try:
            row = self.service.add_product(
                self.stock_name, barcode, product, supplier, measure, mode, critical
            )
        except ValueError as exc:
            messagebox.showerror(APP_TITLE, str(exc))
            return
        messagebox.showinfo(
            APP_TITLE,
            f"Produto cadastrado com sucesso na linha {row}.",
        )
        self.destroy()


# ---------------------------------------------------------------------------
# Account management dialogs
# ---------------------------------------------------------------------------


class ManagerPasswordDialog(simpledialog.Dialog):
    def __init__(self, parent: tk.Misc, current_email: str):
        self.current_email = current_email
        self.result: Optional[Dict[str, str]] = None
        super().__init__(parent, title="Redefinir senha do gestor")

    def body(self, master):
        ttk.Label(master, text="Senha atual:").grid(row=0, column=0, padx=10, pady=(10, 0), sticky="w")
        self.current_entry = ttk.Entry(master, show="*")
        self.current_entry.grid(row=1, column=0, padx=10, pady=(0, 5), sticky="we")

        ttk.Label(master, text="Nova senha (apenas números):").grid(
            row=2, column=0, padx=10, pady=(5, 0), sticky="w"
        )
        self.new_entry = ttk.Entry(master, show="*")
        self.new_entry.grid(row=3, column=0, padx=10, pady=(0, 5), sticky="we")

        ttk.Label(master, text="Confirme a nova senha:").grid(
            row=4, column=0, padx=10, pady=(5, 0), sticky="w"
        )
        self.confirm_entry = ttk.Entry(master, show="*")
        self.confirm_entry.grid(row=5, column=0, padx=10, pady=(0, 5), sticky="we")

        ttk.Label(master, text="E-mail do gestor:").grid(
            row=6, column=0, padx=10, pady=(5, 0), sticky="w"
        )
        self.email_entry = ttk.Entry(master)
        self.email_entry.insert(0, self.current_email)
        self.email_entry.grid(row=7, column=0, padx=10, pady=(0, 10), sticky="we")

        master.columnconfigure(0, weight=1)
        return self.current_entry

    def validate(self) -> bool:
        current = self.current_entry.get().strip()
        new_pin = self.new_entry.get().strip()
        confirm = self.confirm_entry.get().strip()
        email = self.email_entry.get().strip()
        if not current:
            messagebox.showerror(APP_TITLE, "Informe a senha atual.")
            return False
        if not new_pin or not confirm:
            messagebox.showerror(APP_TITLE, "Informe a nova senha nos dois campos.")
            return False
        if new_pin != confirm:
            messagebox.showerror(APP_TITLE, "As senhas novas não coincidem.")
            return False
        if not new_pin.isdigit():
            messagebox.showerror(APP_TITLE, "A nova senha deve conter apenas números.")
            return False
        if len(new_pin) > 10:
            messagebox.showerror(APP_TITLE, "A nova senha deve ter no máximo 10 dígitos.")
            return False
        if not email:
            messagebox.showerror(APP_TITLE, "Informe o e-mail do gestor.")
            return False
        self.result = {
            "current": current,
            "new": new_pin,
            "email": email,
        }
        return True


class StockPasswordDialog(simpledialog.Dialog):
    def __init__(self, parent: tk.Misc):
        self.result: Optional[str] = None
        super().__init__(parent, title="Senha das planilhas")

    def body(self, master):
        ttk.Label(master, text="Nova senha das planilhas:").grid(
            row=0, column=0, padx=10, pady=(10, 0), sticky="w"
        )
        self.new_entry = ttk.Entry(master, show="*")
        self.new_entry.grid(row=1, column=0, padx=10, pady=(0, 5), sticky="we")

        ttk.Label(master, text="Confirme a nova senha:").grid(
            row=2, column=0, padx=10, pady=(5, 0), sticky="w"
        )
        self.confirm_entry = ttk.Entry(master, show="*")
        self.confirm_entry.grid(row=3, column=0, padx=10, pady=(0, 10), sticky="we")

        master.columnconfigure(0, weight=1)
        return self.new_entry

    def validate(self) -> bool:
        password = self.new_entry.get().strip()
        confirm = self.confirm_entry.get().strip()
        if not password:
            messagebox.showerror(APP_TITLE, "Informe a nova senha.")
            return False
        if password != confirm:
            messagebox.showerror(APP_TITLE, "As senhas informadas não coincidem.")
            return False
        self.result = password
        return True


# ---------------------------------------------------------------------------
# Branding helpers
# ---------------------------------------------------------------------------


class LogoBanner(tk.Frame):
    """Displays the provided corporate logo image."""

    _cached_image: Optional[tk.PhotoImage] = None

    def __init__(self, master):
        super().__init__(master, bg="white")
        self.columnconfigure(0, weight=1)

        self.label = tk.Label(self, bg="white")
        self.label.grid(row=0, column=0)

        image = self._load_image()
        if image is not None:
            self.label.configure(image=image)
            self.label.image = image  # prevent garbage collection
        else:
            self.label.configure(
                text="LOGO.png não encontrado",
                font=("Segoe UI", 14, "bold"),
                fg="#b91d1d",
            )

    @classmethod
    def _load_image(cls) -> Optional[tk.PhotoImage]:
        if cls._cached_image is not None:
            return cls._cached_image

        if not LOGO_FILE.exists():
            return None

        try:
            cls._cached_image = tk.PhotoImage(file=str(LOGO_FILE))
        except tk.TclError:
            cls._cached_image = None
        return cls._cached_image


# ---------------------------------------------------------------------------
# Main application controller
# ---------------------------------------------------------------------------


class InventoryApp:
    def __init__(self, root: tk.Tk, service: InventoryService):
        self.root = root
        self.service = service
        self.manager_mode = False
        self.sheet_preview_tree: Optional[ttk.Treeview] = None
        self.sheet_selection_vars: Dict[str, tk.BooleanVar] = {}
        self.root.title(APP_TITLE)
        self.root.resizable(True, True)
        self.root.configure(bg="white")
        self.root.geometry("1200x700")

        self._build_public_interface()

    # -- interface builders ---------------------------------------------

    def _clear_window(self):
        for widget in self.root.winfo_children():
            widget.destroy()

    def _build_public_interface(self):
        self.manager_mode = False
        self._clear_window()
        frame = tk.Frame(self.root, bg="white")
        frame.pack(fill="both", expand=True, padx=20, pady=20)

        toolbar = self._build_toolbar(frame)
        ttk.Button(
            toolbar, text="Registrar saída", command=self._public_register_exit
        ).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Abrir scanner", command=self._public_scanner).pack(
            side="left", padx=5
        )
        ttk.Button(toolbar, text="Recuperar senha", command=self._recover_password).pack(
            side="left", padx=5
        )
        tk.Button(
            toolbar,
            text="Login",
            command=self._prompt_login,
            bg="#241178",
            fg="white",
            activebackground="#3620a0",
            activeforeground="white",
        ).pack(side="right", padx=5)

        LogoBanner(frame).pack(pady=(0, 10))

        tk.Label(
            frame,
            text="Controle de Estoque",
            font=("Segoe UI", 14, "bold"),
            bg="white",
        ).pack(pady=(0, 10))

        self._build_sheet_preview(frame)

    def _build_manager_interface(self):
        self.manager_mode = True
        self._clear_window()
        frame = tk.Frame(self.root, bg="white")
        frame.pack(fill="both", expand=True, padx=20, pady=20)

        toolbar = self._build_toolbar(frame)
        ttk.Button(
            toolbar,
            text="Registrar entrada",
            command=lambda: self._open_manual_movement(is_entry=True),
        ).pack(side="left", padx=5)
        ttk.Button(
            toolbar,
            text="Registrar saída",
            command=lambda: self._open_manual_movement(is_entry=False),
        ).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Abrir scanner", command=self._manager_scanner).pack(
            side="left", padx=5
        )
        ttk.Button(toolbar, text="Novo produto", command=self._new_product).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Abrir planilha", command=self._open_workbook_file).pack(
            side="left", padx=5
        )
        ttk.Button(toolbar, text="Redefinir senha", command=self._reset_manager_password).pack(
            side="left", padx=5
        )
        ttk.Button(
            toolbar, text="Senha das planilhas", command=self._update_stock_password
        ).pack(side="left", padx=5)
        tk.Button(
            toolbar,
            text="Logout",
            command=self._build_public_interface,
            bg="#241178",
            fg="white",
            activebackground="#3620a0",
            activeforeground="white",
        ).pack(side="right", padx=5)

        LogoBanner(frame).pack(pady=(0, 10))

        tk.Label(
            frame,
            text="Controle de Estoque",
            font=("Segoe UI", 14, "bold"),
            bg="white",
        ).pack(pady=(0, 10))

        self._build_sheet_preview(frame)

    def _build_toolbar(self, parent: tk.Misc) -> tk.Frame:
        toolbar = tk.Frame(parent, bg="white")
        toolbar.pack(fill="x", pady=(0, 10))
        return toolbar

    def _build_sheet_preview(self, parent: tk.Misc) -> None:
        selector = tk.Frame(parent, bg="white")
        selector.pack(fill="x", pady=(0, 8))

        tk.Label(
            selector,
            text="Visualizar planilha:",
            bg="white",
            font=("Segoe UI", 10, "bold"),
        ).pack(side="left", padx=(0, 8))

        self.sheet_selection_vars = {}
        for name in self.service.list_all_planilhas():
            var = tk.BooleanVar(value=False)
            self.sheet_selection_vars[name] = var
            ttk.Checkbutton(
                selector,
                text=name,
                variable=var,
                command=lambda sheet_name=name: self._on_sheet_toggle(sheet_name),
            ).pack(side="left", padx=4)

        preview_frame = tk.Frame(
            parent,
            bg="white",
            highlightbackground="#4a4a4a",
            highlightthickness=1,
        )
        preview_frame.pack(fill="both", expand=True)
        preview_frame.columnconfigure(0, weight=1)
        preview_frame.rowconfigure(0, weight=1)

        style = ttk.Style()
        style.configure("SheetPreview.Treeview", rowheight=24)
        style.configure(
            "SheetPreview.Treeview.Heading",
            background="#FFCC80",
            foreground="black",
        )

        tree = ttk.Treeview(preview_frame, columns=(), show="headings", style="SheetPreview.Treeview")
        tree.grid(row=0, column=0, sticky="nsew")
        self.sheet_preview_tree = tree

        tree.tag_configure("headers", background="#FFE0B2")
        tree.tag_configure("even", background="white")
        tree.tag_configure("odd", background="#FFF3E0")

        y_scroll = ttk.Scrollbar(preview_frame, orient="vertical", command=tree.yview)
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll = ttk.Scrollbar(preview_frame, orient="horizontal", command=tree.xview)
        x_scroll.grid(row=1, column=0, sticky="ew")
        tree.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)

    def _on_sheet_toggle(self, sheet_name: str) -> None:
        var = self.sheet_selection_vars.get(sheet_name)
        if var is None:
            return
        if var.get():
            for other_name, other_var in self.sheet_selection_vars.items():
                if other_name != sheet_name:
                    other_var.set(False)
            self._load_sheet_preview(sheet_name)
        else:
            self._clear_sheet_preview()

    def _clear_sheet_preview(self) -> None:
        if not self.sheet_preview_tree:
            return
        for item in self.sheet_preview_tree.get_children():
            self.sheet_preview_tree.delete(item)
        self.sheet_preview_tree.configure(columns=())

    def _load_sheet_preview(self, sheet_name: str) -> None:
        if not self.sheet_preview_tree:
            return
        try:
            columns, rows = self.service.get_sheet_preview(sheet_name)
        except ValueError as exc:
            messagebox.showerror(APP_TITLE, str(exc))
            self._clear_sheet_preview()
            return
        self.sheet_preview_tree.configure(columns=columns)
        for col in columns:
            self.sheet_preview_tree.heading(col, text=col)
            self.sheet_preview_tree.column(col, width=140, anchor="w")
        self.sheet_preview_tree.delete(*self.sheet_preview_tree.get_children())
        for values, tag in rows:
            self.sheet_preview_tree.insert("", "end", values=values, tags=(tag,))

    # -- event handlers --------------------------------------------------

    def _select_workbook(
        self,
        *,
        include_log: bool = False,
        title: str = "Saída – escolha o setor",
        prompt: str = "Escolha o estoque:",
    ) -> Optional[str]:
        choices = (
            self.service.list_all_planilhas() if include_log else self.service.list_stock_names()
        )
        selector = WorkbookSelector(self.root, choices, title=title, prompt=prompt)
        return selector.selection

    def _select_stock(self) -> Optional[str]:
        return self._select_workbook()

    def _open_manual_movement(self, is_entry: bool):
        stock_name = self._select_stock()
        if not stock_name:
            return
        ManualMovementWindow(
            self.root, self.service, stock_name, is_entry=is_entry, manager_mode=self.manager_mode
        )

    def _public_register_exit(self):
        self._open_manual_movement(is_entry=False)

    def _public_scanner(self):
        stock_name = self._select_stock()
        if not stock_name:
            return
        ScannerWindow(self.root, self.service, stock_name, manager_mode=False)

    def _manager_scanner(self):
        stock_name = self._select_stock()
        if not stock_name:
            return
        ScannerWindow(self.root, self.service, stock_name, manager_mode=True)

    def _new_product(self):
        stock_name = self._select_stock()
        if not stock_name:
            return
        NewProductWindow(self.root, self.service, stock_name)

    def _prompt_login(self):
        pin = simpledialog.askstring("Acesso restrito", "PIN do Gestor:", show="*")
        if pin is None:
            return
        if self.service.verify_manager_pin(pin):
            messagebox.showinfo(APP_TITLE, "Acesso liberado.")
            self._build_manager_interface()
        else:
            messagebox.showerror(APP_TITLE, "PIN incorreto.")

    def _recover_password(self):
        if not messagebox.askyesno(
            APP_TITLE,
            "Deseja receber a senha cadastrada no e-mail informado?",
        ):
            return
        try:
            self.service.recover_manager_pin()
        except ValueError as exc:
            messagebox.showerror(APP_TITLE, str(exc))
            return
        email = self.service.get_manager_email()
        messagebox.showinfo(
            APP_TITLE,
            f"A senha cadastrada foi enviada para o e-mail {email}.",
        )

    def _reset_manager_password(self):
        dialog = ManagerPasswordDialog(self.root, self.service.get_manager_email())
        if not dialog.result:
            return
        try:
            self.service.update_manager_pin(
                dialog.result["current"], dialog.result["new"], dialog.result["email"]
            )
        except ValueError as exc:
            messagebox.showerror(APP_TITLE, str(exc))
            return
        email = dialog.result["email"].strip()
        messagebox.showinfo(
            APP_TITLE,
            f"Senha atualizada com sucesso. A nova senha foi enviada para o e-mail {email}.",
        )

    def _update_stock_password(self):
        dialog = StockPasswordDialog(self.root)
        if not dialog.result:
            return
        try:
            self.service.update_stock_password(dialog.result)
        except ValueError as exc:
            messagebox.showerror(APP_TITLE, str(exc))
            return
        messagebox.showinfo(APP_TITLE, "Senha das planilhas atualizada com sucesso.")

    def _open_workbook_file(self):
        selection = self._select_workbook(
            include_log=True,
            title="Abrir planilha",
            prompt="Escolha a planilha:",
        )
        if not selection:
            return
        try:
            path = self.service.get_workbook_path(selection)
            open_file_with_default_app(path)
        except ValueError as exc:
            messagebox.showerror(APP_TITLE, str(exc))


# ---------------------------------------------------------------------------
# Entrypoint
# ---------------------------------------------------------------------------


def main() -> None:
    try:
        settings = AppSettings()
    except RuntimeError as exc:
        temp_root = tk.Tk()
        temp_root.withdraw()
        messagebox.showerror(APP_TITLE, str(exc), parent=temp_root)
        temp_root.destroy()
        return
    email_service = EmailService(settings)
    service = InventoryService(settings, email_service)
    root = tk.Tk()
    InventoryApp(root, service)
    root.mainloop()


if __name__ == "__main__":
    main()
